"""
Verification gate for AITHON_2.0_RetinaScanAI_Submission.pptx  (v2 — chrome-only).

Template rigidity level: **Relaxed — chrome only.**  Approved by the submitter:
template *content* placeholder shapes may be moved, deleted or re-texted freely;
only the items below are verified byte/geometry-identical.

Checks (each prints PASS/FAIL; process exits non-zero on any FAIL):
   1.  Template file untouched (sha256 of AITHON_2.0_Presentation.pptx)
   2.  Slide count = 7, slide size = 10.0 x 5.625 in (both files)
   3.  Section headings byte-identical to template (named heading shapes)
   4.  Footer bar text + geometry identical on every slide
   5.  Both template logos present on every slide with identical
       name + geometry (additional images — charts/screenshots — are allowed)
   6.  Slide 7 (IMPORTANT INSTRUCTIONS) text content identical to template
   7.  Placeholder scan: no leftover template prompts; bracketed placeholders
       only the 2 intentional ones on slide 1
   8.  All shapes within slide bounds (+0.06 in tolerance for template quirks)
   9.  All latin fonts within the allowed set (template fonts + Georgia)
  10.  Every numeric token in the deck traceable to a documented source
       (PROJECT_INFO.md / METRICS.md / hackathon_sprint/metrics/*.json /
       models/*/training_summary.json / models/onnx/benchmark_results.json /
       template chrome)
  11.  Speaker notes present on slides 1-6; slide 7 notes unchanged
  12.  docProps title/author set

Removed vs v1 (per relaxed gate): the v1 "every template content shape still
present with identical geometry" check and its builder-prefix guard.
"""

import hashlib
import re
import sys
from pathlib import Path

import lxml.etree as etree
from pptx import Presentation
from pptx.util import Emu

PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = PROJECT_ROOT / "AITHON_2.0_Presentation.pptx"
OUTPUT = PROJECT_ROOT / "AITHON_2.0_RetinaScanAI_Submission.pptx"

TEMPLATE_SHA256 = "e57ec40b57f0a5785c3f0bd2ffd503f57b220a289f4bb04e7b1076e4880e8350"
FOOTER_TEXT = "@AITHON2.0 Template Submission"
SLIDE_W_EMU, SLIDE_H_EMU = 9144000, 5143500  # 10.0 x 5.625 in
TOL_IN = 0.06

HEADINGS = {
    1: {"Text 21": "AITHON 2.0", "Text 22": "TITLE PAGE"},
    2: {"Text 22": "IDEA TITLE"},
    3: {"Text 22": "TECHNICAL APPROACH"},
    4: {"Text 22": "FEASIBILITY AND VIABILITY"},
    5: {"Text 22": "IMPACT AND BENEFITS"},
    6: {"Text 22": "RESEARCH AND REFERENCES"},
    7: {"Text 22": "IMPORTANT INSTRUCTIONS"},
}

# Logo pictures are read per slide from the template (their shape names vary:
# s1 Picture 46/4, s2 Picture 49/4, s3 Picture 51/4, s4-s7 Picture 65/2 ...).
# Each template slide carries exactly two pictures — the AIESA logo top-left
# and the AITHON 2.0 logo top-right — and both must survive with identical
# name + geometry.  Extra images in the output (charts/screenshots) are allowed.
LOGO_TOL_IN = 0.005
EXPECTED_LOGOS_PER_SLIDE = 2

FORBIDDEN_PROMPTS = [
    "[ Enter the project or idea name ]",
    "[ A crisp single-line tagline that captures the idea ]",
    "[ Describe the core problem this idea solves in 2–3 lines ]",
    "[ Summarize the proposed solution and its key differentiators ]",
    "Frontend & Backend frameworks",
    "AI / ML libraries & models",
    "Cloud & deployment tools",
    "Version control & CI/CD",
    "Modular, layered architecture",
    "AI/ML components for core logic",
    "Database, API & cloud integration",
    "Secure, scalable request flow",
    "Availability of required technologies",
    "Development requirements & resources",
    "Ease of implementation",
    "Ease of adoption by users",
    "Development cost",
    "Infrastructure / resource requirements",
    "Future expansion potential",
    "Long-term usability, maintenance & scalability",
    "User Benefits",
    "Social Impact",
    "Economic Impact",
    "Environmental Impact",
    "Efficiency & Productivity",
    "Relevant research papers",
    "Existing solutions & products",
    "Industry reports & market studies",
    "Datasets used for training / testing",
    "Research papers (IEEE / ACM / arXiv, etc.)",
    "Official websites & documentation",
    "APIs / SDKs referenced",
    "Government / official data sources",
]
ALLOWED_BRACKETS = {"[ Enter track name ]", "[ Enter team name ]"}

ALLOWED_FONTS = {
    "Times New Roman",
    "Calibri",
    "Georgia",
    "Cambria",
    "Microsoft Yi Baiti",
    "+mj-lt",
}

# Every numeric token allowed in slide text, with its documented source.
ALLOWED_NUMBERS = {
    # --- template chrome -------------------------------------------------
    "2.0": "template chrome (AITHON 2.0 heading / footer)",
    # --- template slide-7 instructions wording ---------------------------
    "1": "template slide 7 instruction text",
    "2": "template slide 7 instruction text",
    "3": "template slide 7 instruction text",
    "4": "FHIR R4 / 4 lesion masks / ICDR stage 0-4 (PROJECT_INFO)",
    "4,": "'FHIR R4,' (PROJECT_INFO)",
    "5": "ICDR 5-stage grading (PROJECT_INFO)",
    "7": "'HL7' (PROJECT_INFO)",
    "7.": "'hl7.org' (PROJECT_INFO references)",
    "10": "PROJECT_INFO: '5-10 min' manual grading / batch triage of 10",
    "11": "PROJECT_INFO: all 11 components",
    "01": "pipeline step 1 (builder flow labels)",
    "02": "pipeline step 2 (builder flow labels)",
    "03": "pipeline step 3 (builder flow labels)",
    "04": "pipeline step 4 (builder flow labels)",
    "05": "pipeline step 5 (builder flow labels)",
    "06": "pipeline step 6 (builder flow labels)",
    "07": "pipeline step 7 (builder flow labels)",
    "6": "reference list item 6 (slide 6 citations)",
    # --- cost / time / reach --------------------------------------------
    "15": "PROJECT_INFO: $15-50 per manual screening",
    "50": "PROJECT_INFO: $15-50 per manual screening",
    "77": "PROJECT_INFO: '~77 million diabetics' at risk in India",
    "0": "PROJECT_INFO: $0 per screening after setup",
    "100": "PROJECT_INFO: 100% local / 100x faster",
    "20": "PROJECT_INFO: batch triage of 20 images per call",
    "55": "PROJECT_INFO: abstention below 55% confidence",
    # --- runtime / stack --------------------------------------------------
    "0.59": "PROJECT_INFO: full pipeline 0.59 s",
    "0.71": "PROJECT_INFO: P95 0.71 s",
    "95": "PROJECT_INFO: P95 0.71 s / 95% confidence interval",
    "3.12": "PROJECT_INFO: Python 3.12",
    "2.6": "PROJECT_INFO: PyTorch 2.6",
    "0.141": "PROJECT_INFO: FastAPI 0.141 backend",
    "19": "PROJECT_INFO: React 19",
    "18": "PROJECT_INFO: U-Net + ResNet18 encoder (segmenter)",
    "32": "PROJECT_INFO: ONNX Runtime FP32 primary precision",
    "28": "PROJECT_INFO: ConvNeXt-Tiny 28M params",
    "12.5": "PROJECT_INFO: 12.5M segmenter params",
    "512": "PROJECT_INFO: 512^2 preprocessing crop",
    "3000": "PROJECT_INFO: frontend :3000 (start_demo.sh)",
    "8000": "PROJECT_INFO: backend :8000 (start_demo.sh)",
    # --- live-run evidence (documented generated artifact) ----------------
    "52.1": "scripts/deck_assets/out/evidence.json confidence=0.5206 (capture_evidence.py)",
    # --- headline results (metrics JSON) ----------------------------------
    "82.1": "metrics/ensemble_results.json ensemble_acc=0.8205",
    "97.5": "metrics/external_validation.json sens=0.9748 npv=0.9755",
    "96.5": "metrics/external_validation.json sens 95% CI lower 0.9654",
    "98.2": "metrics/external_validation.json sens 95% CI upper 0.9817",
    "3,394": "metrics/external_validation.json evaluated=3394",
    "3,662": "metrics/external_validation.json total_images=3662",
    "7.3": "metrics/external_validation.json rejection_rate=0.0732",
    "60.8": "metrics/external_validation.json accuracy=0.6081",
    "59.2": "metrics/external_validation.json accuracy 95% CI lower",
    "62.3": "metrics/external_validation.json accuracy 95% CI upper",
    "0.843": "metrics/external_validation.json referable_dr.auc=0.8433",
    "72.8": "metrics/external_validation.json referable_dr.specificity=0.7284",
    "72.3": "metrics/external_validation.json referable_dr.ppv=0.7227",
    "1,392": "metrics/external_validation.json referable_dr.tp=1392",
    "534": "metrics/external_validation.json referable_dr.fp=534",
    "1,432": "metrics/external_validation.json referable_dr.tn=1432",
    "36": "metrics/external_validation.json referable_dr.fn=36",
    "350": "METRICS.md: classifier inference mean 350 ms",
    "445": "METRICS.md: classifier inference p95 445 ms",
    "489": "METRICS.md: classifier inference p99 489 ms",
    "0.314": "models/segmentation/training_summary.json test_dice=0.3144",
    "0.769": "models/segmentation/training_summary.json cotton_wool_spots=0.7692",
    "54": "PROJECT_INFO / METRICS.md: segmenter 54 -> 811 images",
    "811": "PROJECT_INFO: segmenter 811 masked images",
    "413": "METRICS.md: classifier training 413 -> 4,075 images",
    "4,075": "METRICS.md: classifier training 413 -> 4,075 images",
    "815": "models/classification/classification_report.txt support=815",
    "78.4": "models/classification/classification_report.txt accuracy=0.7840",
    "73.5": "PROJECT_INFO: EfficientNet-B2 baseline val accuracy 73.5%",
    "80.4": "PROJECT_INFO / METRICS.md: best ConvNeXt checkpoint 80.4%",
    "8.6": "PROJECT_INFO: +8.6 pp over EfficientNet-B2 baseline",
    "1.7": "derived: 82.1% - 80.4% = 1.7 pp (ensemble over single best)",
    "225": "models/segmentation/training_summary.json DDR test split=225",
    # --- dates / citations -------------------------------------------------
    "2026": "project meta: October 2026",
    "2015": "PROJECT_INFO: U-Net MICCAI 2015",
    "2017": "PROJECT_INFO: Grad-CAM / Focal Loss ICCV 2017",
    "2018": "PROJECT_INFO: IDRiD (Porwal et al., Data 2018)",
    "2019": "PROJECT_INFO: APTOS 2019",
    "2020": "PROJECT_INFO: 'A ConvNet for the 2020s'",
    "2022": "PROJECT_INFO: ConvNeXt CVPR 2022",
    "1505.04597": "PROJECT_INFO: U-Net arXiv id",
    "1610.02391": "PROJECT_INFO: Grad-CAM arXiv id",
    "2201.03545": "PROJECT_INFO: ConvNeXt arXiv id",
}

results = []


def check(label, ok, detail=""):
    results.append((label, bool(ok), detail))
    print(
        f"{'PASS' if ok else 'FAIL'}  {label}"
        + (f"  — {detail}" if detail and not ok else "")
    )


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def geo(shape):
    return (shape.left, shape.top, shape.width, shape.height)


def geo_in(shape):
    return (
        round(Emu(shape.left).inches, 3),
        round(Emu(shape.top).inches, 3),
        round(Emu(shape.width).inches, 3),
        round(Emu(shape.height).inches, 3),
    )


def main():
    # 1. Template untouched
    check(
        "1. template sha256 unchanged",
        sha256(TEMPLATE) == TEMPLATE_SHA256,
        f"got {sha256(TEMPLATE)}",
    )
    tpl = Presentation(str(TEMPLATE))
    out = Presentation(str(OUTPUT))

    # 2. slide count + size
    check(
        "2a. slide count = 7",
        len(out.slides) == 7 and len(tpl.slides) == 7,
        f"out={len(out.slides)} tpl={len(tpl.slides)}",
    )
    sizes_ok = (out.slide_width, out.slide_height) == (SLIDE_W_EMU, SLIDE_H_EMU) and (
        tpl.slide_width,
        tpl.slide_height,
    ) == (SLIDE_W_EMU, SLIDE_H_EMU)
    check(
        "2b. slide size 10.0 x 5.625 in",
        sizes_ok,
        f"out={out.slide_width}x{out.slide_height}",
    )

    all_fonts, all_numbers, bracket_texts = set(), set(), set()
    extra_pics = []

    for idx in range(1, 8):
        ts, os_ = tpl.slides[idx - 1], out.slides[idx - 1]
        tmap, omap = (
            {sh.name: sh for sh in ts.shapes},
            {sh.name: sh for sh in os_.shapes},
        )

        # 3. headings
        for name, expected in HEADINGS[idx].items():
            got = omap[name].text_frame.text if name in omap else "<missing>"
            check(f"3. s{idx} heading {name!r}", got == expected, f"got {got!r}")

        # 4. footer text + geometry
        f_txt = [
            sh
            for sh in os_.shapes
            if sh.has_text_frame and sh.text_frame.text == FOOTER_TEXT
        ]
        t_foot = [
            sh
            for sh in ts.shapes
            if sh.has_text_frame and sh.text_frame.text == FOOTER_TEXT
        ]
        tf_ok = len(f_txt) == 1 and len(t_foot) == 1 and geo(f_txt[0]) == geo(t_foot[0])
        check(
            f"4. s{idx} footer text+geometry",
            tf_ok,
            f"out={len(f_txt)} tpl={len(t_foot)}",
        )

        # 5. both template logos present with identical geometry
        #    (v2: logo names are taken from the template slide; additional
        #     chart/screenshot images in the output are permitted)
        t_pics = {sh.name: geo_in(sh) for sh in ts.shapes if sh.shape_type == 13}
        logo_problems = []
        if len(t_pics) != EXPECTED_LOGOS_PER_SLIDE:
            logo_problems.append(
                f"template has {len(t_pics)} pictures, expected {EXPECTED_LOGOS_PER_SLIDE}"
            )
        for lname, want in t_pics.items():
            if lname not in omap:
                logo_problems.append(f"{lname} missing")
            elif any(
                abs(a - b) > LOGO_TOL_IN for a, b in zip(geo_in(omap[lname]), want)
            ):
                logo_problems.append(f"{lname} moved {geo_in(omap[lname])} != {want}")
        check(
            f"5. s{idx} both logos present, geometry identical",
            not logo_problems,
            str(logo_problems),
        )
        for sh in os_.shapes:
            if sh.shape_type == 13 and sh.name not in t_pics:
                extra_pics.append(f"s{idx}:{sh.name}")

        # 6. slide 7 verbatim (content untouched under the relaxed gate)
        if idx == 7:
            t_txt = [sh.text_frame.text for sh in ts.shapes if sh.has_text_frame]
            o_txt = [sh.text_frame.text for sh in os_.shapes if sh.has_text_frame]
            check(
                "6. slide 7 verbatim vs template",
                t_txt == o_txt,
                f"diff: {set(t_txt) ^ set(o_txt)}",
            )
            t_geo = sorted(geo(sh) for sh in ts.shapes)
            o_geo = sorted(geo(sh) for sh in os_.shapes)
            check(
                "6b. slide 7 shapes+geometry identical",
                t_geo == o_geo,
                "geometry differs",
            )

        # 7. placeholders + numbers + fonts
        for sh in os_.shapes:
            if not sh.has_text_frame:
                continue
            txt = sh.text_frame.text
            if idx in (2, 3, 4, 5, 6):
                for p in FORBIDDEN_PROMPTS:
                    if txt == p:
                        check(
                            f"7. s{idx} no leftover prompt {p[:30]!r}",
                            False,
                            "still present",
                        )
            bracket_texts.update(re.findall(r"\[[^\]]*\]", txt))
            all_numbers.update(re.findall(r"\d[\d,.]*", txt))
        # fonts (latin typefaces in raw slide xml)
        root = etree.fromstring(os_._element.xml.encode())
        ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}latin"
        for el in root.iter(ns):
            if el.get("typeface"):
                all_fonts.add(el.get("typeface"))

    # 7b. bracket placeholders only the two intentional ones
    check(
        "7b. bracketed placeholders = only Track + Team Name",
        bracket_texts <= ALLOWED_BRACKETS,
        f"unexpected={sorted(bracket_texts - ALLOWED_BRACKETS)}",
    )

    # 8. bounds
    oob = []
    for idx, s in enumerate(out.slides, 1):
        for sh in s.shapes:
            l, t = Emu(sh.left).inches, Emu(sh.top).inches
            r, b = l + Emu(sh.width).inches, t + Emu(sh.height).inches
            if (
                l < -TOL_IN
                or t < -TOL_IN - 0.01
                or r > 10.0 + TOL_IN
                or b > 5.625 + TOL_IN
            ):
                oob.append(f"s{idx}:{sh.name}({l:.2f},{t:.2f})-({r:.2f},{b:.2f})")
    check("8. all shapes within slide bounds", not oob, f"out-of-bounds={oob}")

    # 9. fonts
    bad_fonts = all_fonts - ALLOWED_FONTS
    check(
        "9. fonts within allowed set",
        not bad_fonts,
        f"disallowed={sorted(bad_fonts)} (all={sorted(all_fonts)})",
    )

    # 10. numbers traceable
    unknown = sorted(all_numbers - set(ALLOWED_NUMBERS))
    check(
        "10. every numeric token sourced",
        not unknown,
        f"unsourced={unknown} (all tokens={sorted(all_numbers)})",
    )

    # 11. notes
    notes_ok, n_detail = True, ""
    for idx in range(1, 7):
        n = out.slides[idx - 1].notes_slide.notes_text_frame.text.strip()
        if len(n) < 100:
            notes_ok, n_detail = False, f"slide {idx} notes too short ({len(n)})"
    if (
        tpl.slides[6].notes_slide.notes_text_frame.text
        != out.slides[6].notes_slide.notes_text_frame.text
    ):
        notes_ok, n_detail = False, "slide 7 notes changed"
    check("11. notes present s1-s6, s7 unchanged", notes_ok, n_detail)

    # 12. docProps
    cp = out.core_properties
    check(
        "12. docProps title/author",
        (cp.title or "").startswith("AITHON 2.0") and (cp.author or "") == "Siddhesh",
        f"title={cp.title!r} author={cp.author!r}",
    )

    fails = [r for r in results if not r[1]]
    print("\n" + "=" * 60)
    print(f"RESULT: {len(results) - len(fails)}/{len(results)} checks passed")
    if extra_pics:
        print(
            f"additional images (charts/screenshots, allowed): {', '.join(extra_pics)}"
        )
    if fails:
        print("FAILED:")
        for label, _, detail in fails:
            print(f"  - {label}: {detail}")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
