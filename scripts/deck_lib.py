"""
Shared design-system primitives for the redesigned AITHON deck.

Every helper writes template-safe shapes: text boxes with explicit sizing
(never autofit), cards with hairline borders, KPI blocks, chips, flow boxes,
arrows and circular photo crops.  Palette and type scale are documented in
scripts/deck_assets/DESIGN.md.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ── Palette (template-derived) ──────────────────────────────────────────────
ACCENT = RGBColor(0x00, 0x70, 0xC0)  # template footer blue
DEEP = RGBColor(0x1E, 0x27, 0x61)  # template navy
VIOLET = RGBColor(0x6C, 0x63, 0xFF)  # template violet
INK = RGBColor(0x10, 0x18, 0x28)
SLATE = RGBColor(0x47, 0x55, 0x69)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
LINE = RGBColor(0xE2, 0xE8, 0xF0)
PANEL = RGBColor(0xF8, 0xFA, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
POS = RGBColor(0x0E, 0x9F, 0x9F)
WARN = RGBColor(0xD9, 0x77, 0x06)
NEG = RGBColor(0xDC, 0x26, 0x26)
SOFT = RGBColor(0xFE, 0xF2, 0xF2)  # alert card fill

FONT = "Calibri"
DISPLAY = "Georgia"

# ── Geometry ────────────────────────────────────────────────────────────────
SLIDE_W, SLIDE_H = 10.0, 5.625
MARGIN_L, MARGIN_R = 0.55, 9.45
CONTENT_W = MARGIN_R - MARGIN_L  # 8.90
FOOTER_Y = 5.32


# ── low-level ───────────────────────────────────────────────────────────────
def _no_autofit(tf) -> None:
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)


def _style_run(run, spec: dict) -> None:
    f = run.font
    f.name = spec.get("font", FONT)
    f.size = Pt(spec.get("size", 9))
    f.bold = spec.get("bold", False)
    f.italic = spec.get("italic", False)
    f.color.rgb = spec.get("color", SLATE)
    if "spc" in spec:  # letter-spacing, hundredths of a point
        run.font._rPr.set("spc", str(int(spec["spc"])))
    if spec.get("link"):  # real clickable hyperlink (text runs only)
        run.hyperlink.address = spec["link"]
        f.underline = spec.get("underline", False)
        f.color.rgb = spec.get("color", SLATE)  # keep explicit colour over theme


def add_text(
    slide,
    name: str,
    x: float,
    y: float,
    w: float,
    h: float,
    paras: list[dict],
    anchor=MSO_ANCHOR.TOP,
    align=PP_ALIGN.LEFT,
    margins: float = 0.0,
):
    """Add a text box.  Each para dict: text|runs, size, bold, color, font,
    align, space_before, space_after, line (line spacing multiple)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    tf = box.text_frame
    _no_autofit(tf)
    tf.vertical_anchor = anchor
    m = Inches(margins)
    tf.margin_left = tf.margin_right = m
    tf.margin_top = tf.margin_bottom = Emu(0)

    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", align)
        if "space_before" in spec:
            p.space_before = Pt(spec["space_before"])
        if "space_after" in spec:
            p.space_after = Pt(spec["space_after"])
        if "line" in spec:
            p.line_spacing = spec["line"]
        runs = spec.get("runs")
        if runs is None:
            runs = [(spec.get("text", ""), spec)]
        for text, rspec in runs:
            run = p.add_run()
            run.text = text
            _style_run(run, rspec)
    return box


def retint_links(prs, hex_rgb: str) -> None:
    """Paint hyperlink runs in the deck palette.

    Renderers (LibreOffice at least) ignore an explicit run colour on a
    hyperlink and use the theme's <a:hlink>/<a:folHlink> colour instead, so
    retarget the theme token too — otherwise links come out theme-blue and
    shouty.  Only touches /ppt/theme/*.xml; template file itself is untouched.
    """
    import re

    for part in prs.part.package.iter_parts():
        if not str(part.partname).startswith("/ppt/theme/"):
            continue
        xml = part.blob.decode("utf-8")
        new = re.sub(
            r"(<a:hlink>|<a:folHlink>)<a:srgbClr val=\"[0-9A-Fa-f]{6}\"/>",
            lambda m: f'{m.group(1)}<a:srgbClr val="{hex_rgb}"/>',
            xml,
        )
        if new != xml:
            part._blob = new.encode("utf-8")


def add_rect(
    slide,
    name: str,
    x: float,
    y: float,
    w: float,
    h: float,
    fill=None,
    line=None,
    line_pt: float = 0.75,
    shape=MSO_SHAPE.RECTANGLE,
    adj: float | None = None,
):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = name
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_pt)
    if adj is not None and sp.adjustments:
        sp.adjustments[0] = adj
    tf = sp.text_frame
    _no_autofit(tf)
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    return sp


def card(slide, name, x, y, w, h, fill=PANEL, line=LINE, accent=None, accent_h=0.045):
    """Panel with an optional accent top bar.  Returns the card shape."""
    sp = add_rect(slide, name, x, y, w, h, fill=fill, line=line, line_pt=0.75)
    if accent is not None:
        add_rect(slide, f"{name}Bar", x, y, w, accent_h, fill=accent)
    return sp


def rule(slide, name, x, y, w, color=ACCENT, h=0.024):
    return add_rect(slide, name, x, y, w, h, fill=color)


def panel_title(slide, name, x, y, w, text, color=SLATE, size=7.2, rule_after=False):
    add_text(
        slide,
        name,
        x,
        y,
        w,
        0.16,
        [{"text": text.upper(), "size": size, "bold": True, "color": color, "spc": 55}],
    )
    if rule_after:
        rule(slide, f"{name}Rule", x, y + 0.185, w, color=LINE, h=0.012)


def kpi(slide, name, x, y, w, h, value, label, value_size=20, color=DEEP, sub=None):
    """Big display numeral + caption, accent top rule."""
    add_rect(slide, f"{name}Bar", x, y, w, 0.05, fill=color)
    paras = [
        {
            "text": value,
            "font": DISPLAY,
            "size": value_size,
            "bold": True,
            "color": color,
            "space_before": 6,
        },
        {
            "text": label,
            "size": 7.4,
            "color": SLATE,
            "space_before": 1,
            "line": 0.95,
        },
    ]
    if sub:
        paras.append(
            {
                "text": sub,
                "size": 6.6,
                "color": MUTED,
                "space_before": 1,
            }
        )
    add_text(slide, name, x + 0.10, y + 0.05, w - 0.18, h - 0.08, paras)


def chip(
    slide, name, x, y, w, h, runs, fill=PANEL, line=LINE, size=7.4, align=PP_ALIGN.LEFT
):
    add_rect(slide, name, x, y, w, h, fill=fill, line=line, line_pt=0.6)
    add_text(
        slide,
        f"{name}Txt",
        x + 0.08,
        y,
        w - 0.16,
        h,
        [{"runs": runs, "size": size, "align": align}],
        anchor=MSO_ANCHOR.MIDDLE,
        align=align,
    )


def bullets(
    slide, name, x, y, w, h, items, size=8.6, gap=2.5, color=SLATE, bullet_color=ACCENT
):
    """Dot bullets with accent markers."""
    paras = []
    for i, item in enumerate(items):
        paras.append(
            {
                "runs": [
                    ("• ", {"size": size, "bold": True, "color": bullet_color}),
                    (item, {"size": size, "color": color}),
                ],
                "space_after": gap,
                "line": 1.0,
            }
        )
    return add_text(slide, name, x, y, w, h, paras)


def arrow(slide, name, x, y, w, h, color=ACCENT, direction=MSO_SHAPE.RIGHT_ARROW):
    return add_rect(slide, name, x, y, w, h, fill=color, shape=direction)


def fbox(
    slide,
    name,
    x,
    y,
    w,
    h,
    step,
    title,
    sub,
    accent=ACCENT,
    fill=WHITE,
    title_size=8.4,
    sub_size=6.6,
    dashed=False,
):
    """Flow-chart node: hairline box, step eyebrow, title, sub-caption."""
    sp = add_rect(
        slide, name, x, y, w, h, fill=fill, line=(NEG if dashed else LINE), line_pt=1.0
    )
    if dashed:
        sp.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH
    add_rect(slide, f"{name}Bar", x, y, 0.055, h, fill=(NEG if dashed else accent))
    add_text(
        slide,
        f"{name}Txt",
        x + 0.13,
        y + 0.05,
        w - 0.20,
        h - 0.08,
        [
            {
                "text": step.upper(),
                "size": 5.8,
                "bold": True,
                "color": (NEG if dashed else accent),
                "spc": 50,
            },
            {
                "text": title,
                "size": title_size,
                "bold": True,
                "color": INK,
                "space_before": 1.5,
            },
            {
                "text": sub,
                "size": sub_size,
                "color": MUTED,
                "space_before": 1,
                "line": 0.95,
            },
        ],
    )
    return sp


def foot_note(slide, name, x, y, w, h, text, size=7.2, color=MUTED):
    return add_text(
        slide,
        name,
        x,
        y,
        w,
        h,
        [{"text": text, "size": size, "color": color, "line": 1.0}],
    )


# ── pictures ────────────────────────────────────────────────────────────────
def add_pic(slide, path, name, x, y, w=None, h=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)
    pic.name = name
    return pic


def circle_pic(slide, path, name, x, y, d, ring=True, ring_color=ACCENT, ring_pt=2.0):
    """Insert a picture cropped to a circle (prstGeom override) + accent ring."""
    pic = add_pic(slide, path, name, x, y, w=d, h=d)
    sp_pr = pic._element.spPr
    geom = sp_pr.find(qn("a:prstGeom"))
    if geom is None:
        geom = sp_pr.makeelement(qn("a:prstGeom"), {"prst": "rect"})
        geom.append(sp_pr.makeelement(qn("a:avLst"), {}))
        sp_pr.append(geom)
    geom.set("prst", "ellipse")
    if ring:
        add_rect(
            slide,
            f"{name}Ring",
            x - 0.035,
            y - 0.035,
            d + 0.07,
            d + 0.07,
            fill=None,
            line=ring_color,
            line_pt=ring_pt,
            shape=MSO_SHAPE.OVAL,
        )
    return pic


def square_crop(src: Path, dst: Path, size: int = 760) -> Path:
    """Centre-crop to a square and resize — keeps circular crops undistorted."""
    im = Image.open(src).convert("RGB")
    w, h = im.size
    side = min(w, h)
    im = im.crop(((w - side) // 2, (h - side) // 2, (w + side) // 2, (h + side) // 2))
    im = im.resize((size, size), Image.LANCZOS)
    im.save(dst, quality=92)
    return dst


# ── slide surgery ───────────────────────────────────────────────────────────
def drop_shapes(slide, keep_names: set[str]) -> list[str]:
    """Delete every shape whose name is not in keep_names (chrome stays)."""
    removed = []
    for sh in list(slide.shapes):
        if sh.name not in keep_names:
            removed.append(sh.name)
            sh._element.getparent().remove(sh._element)
    return removed


# heading shapes in the official template (slide 1 has two, slides 2-7 one)
HEADING_NAMES = {"Text 21", "Text 22"}


def chrome_names(slide, template) -> set[str]:
    """Names of template chrome: section heading(s), footer bar, both logos."""
    keep = set()
    for sh in template.shapes:
        if sh.shape_type == 13:  # pictures = logos
            keep.add(sh.name)
        elif sh.name in HEADING_NAMES and sh.has_text_frame:
            keep.add(sh.name)
        elif sh.has_text_frame and sh.text_frame.text.startswith("@AITHON2.0"):
            keep.add(sh.name)
    return keep


def set_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def pic_geom(pic) -> tuple[int, int, int, int]:
    return (pic.left, pic.top, pic.width, pic.height)
